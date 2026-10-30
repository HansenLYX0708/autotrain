"""Open the HawkeyePlus decks in PowerPoint and verify/export the native slides."""
import json
import re
from itertools import product
from pathlib import Path

import pythoncom
import win32com.client
from PIL import Image
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from build import contact_sheet, NAMES, LANGUAGES, COLORS, deck_stem, file_suffix
from model_catalog import load_catalog

HERE = Path(__file__).resolve().parent


def walk(shapes):
    for shape in shapes:
        yield shape
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from walk(shape.shapes)


def com_walk(shapes):
    for i in range(1, shapes.Count+1):
        shape = shapes.Item(i)
        if shape.Type == 6:
            yield from com_walk(shape.GroupItems)
        else:
            yield shape


def main():
    report = []
    catalog = load_catalog()
    pythoncom.CoInitialize()
    app = win32com.client.Dispatch("PowerPoint.Application")
    try:
        for lang, mode in product(LANGUAGES, COLORS):
            path = HERE/f"{deck_stem(lang, mode)}_Editable.pptx"
            prs = Presentation(path)
            assert len(prs.slides) == len(NAMES)
            record = {"file": path.name, "language": lang, "slides": [], "text_overflow": []}
            for index, slide in enumerate(prs.slides):
                objects = list(walk(slide.shapes))
                text = "\n".join(s.text for s in objects if s.has_text_frame)
                assert "SWIFTFLOW" not in text.upper(), "Wrong platform in visible content"
                assert "HAWKEYEPLUS" in text.upper(), "Missing HawkeyePlus branding"
                if lang == "en":
                    assert not re.search(r"[\u4e00-\u9fff]", text), "Chinese visible text in English edition"
                    assert not re.search(r"[\u4e00-\u9fff]", slide.notes_slide.notes_text_frame.text), "Chinese speaker notes"
                if index == 0:
                    keys = ["Dataset", "Model Registry"] if lang == "en" else ["数据集配置", "模型版本发布"]
                    assert all(key in text for key in ["PaddlePaddle", "PyTorch", "anomalib"]+keys)
                for framework in {2: ["PaddleDetection", "TorchDet"], 3: ["PaddleSeg", "TorchSeg"],
                                  4: ["TorchAnomaly"]}.get(index, []):
                    assert all(model["display"] in text for model in catalog[framework]), framework
                if index == 4:
                    assert all(name in text for name in ["ResNet18", "ResNet34", "ResNet50", "ResNet101",
                        "MobileNetV3 Small / Large", "PP-LCNet", "PP-HGNet", "Swin Transformer Tiny"])
                data = {
                    "text_boxes": sum(bool(s.has_text_frame and s.text.strip()) for s in objects),
                    "pictures": sum(s.shape_type == MSO_SHAPE_TYPE.PICTURE for s in objects),
                    "groups": sum(s.shape_type == MSO_SHAPE_TYPE.GROUP for s in objects),
                    "native_objects": sum(s.shape_type != MSO_SHAPE_TYPE.PICTURE for s in objects),
                }
                assert data["text_boxes"] >= 18 and data["pictures"] >= 3 and data["groups"] >= 2, data
                assert not any(s.shape_type == MSO_SHAPE_TYPE.PICTURE and
                               s.width >= prs.slide_width*.9 and s.height >= prs.slide_height*.9
                               for s in objects), "Unexpected flattened slide"
                ids = [n.get("id") for n in slide._element.xpath(".//p:cNvPr")]
                assert len(ids) == len(set(ids)), "Duplicate DrawingML IDs"
                for s in slide.shapes:
                    assert s.left >= -12700 and s.top >= -12700, s.name
                    assert s.left+s.width <= prs.slide_width+12700, s.name
                    assert s.top+s.height <= prs.slide_height+12700, s.name
                record["slides"].append(data)
            presentation = app.Presentations.Open(str(path), ReadOnly=True, Untitled=False, WithWindow=False)
            try:
                for i, name in NAMES:
                    slide = presentation.Slides.Item(i)
                    for shape in com_walk(slide.Shapes):
                        if shape.HasTextFrame and shape.TextFrame.HasText:
                            text = shape.TextFrame2.TextRange
                            if text.BoundWidth > shape.Width+2 or text.BoundHeight > shape.Height+2:
                                record["text_overflow"].append({
                                    "slide": i, "text": text.Text,
                                    "bounds": [text.BoundWidth, text.BoundHeight],
                                    "shape": [shape.Width, shape.Height],
                                })
                    target = HERE/"verification"/f"{i:02d}_{name}_{file_suffix(lang, mode)}.png"
                    slide.Export(str(target), "PNG", 3840, 2160)
                    (HERE/"png"/target.name).write_bytes(target.read_bytes())
                    with Image.open(target) as im:
                        assert im.size == (3840, 2160)
                presentation.SaveAs(str(HERE/f"{deck_stem(lang, mode)}.pdf"), 32)
            finally:
                presentation.Close()
            report.append(record)
    finally:
        # Preserve any presentations the user already had open.
        pythoncom.CoUninitialize()
    (HERE/"verification/report.json").write_text(json.dumps({
        "platform": "HawkeyePlus", "powerpoint_opened": True,
        "slides_exported": len(NAMES)*len(LANGUAGES)*len(COLORS), "decks": report,
    }, ensure_ascii=False, indent=2), encoding="utf-8")
    assert not any(r["text_overflow"] for r in report), report
    contact_sheet()
    print(json.dumps(report, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
