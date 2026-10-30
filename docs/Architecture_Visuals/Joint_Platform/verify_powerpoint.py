"""Validate the joint narrative, editable objects and actual PowerPoint output."""
import json
import re
from itertools import product
from pathlib import Path

import pythoncom
import win32com.client
from PIL import Image
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from build import (PAGES, CHAPTERS, NATIVE_SECTIONS, LANGUAGES, THEMES,
                   deck_stem, suffix, contact_sheets)

HERE = Path(__file__).resolve().parent


def walk(shapes):
    for shape in shapes:
        yield shape
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from walk(shape.shapes)


def com_walk(shapes):
    for i in range(1, shapes.Count+1):
        shape = shapes.Item(i)
        if shape.Type == 6:
            yield from com_walk(shape.GroupItems)
        else:
            yield shape


def main():
    report = []
    pythoncom.CoInitialize()
    app = win32com.client.Dispatch("PowerPoint.Application")
    try:
        for lang, mode in product(LANGUAGES, THEMES):
            path = HERE/f"{deck_stem(lang, mode)}_Editable.pptx"
            prs = Presentation(path)
            assert len(prs.slides) == len(PAGES) == 20
            assert [m["kind"] for m in PAGES[:3]] == ["cover", "agenda", "divider"]
            assert [m["kind"] for m in PAGES[-2:]] == ["summary", "closing"]
            assert sum(m["kind"] == "divider" for m in PAGES) == 4
            record = {"file": path.name, "language": lang, "theme": mode, "slides": [], "text_overflow": []}
            for index, slide in enumerate(prs.slides, 1):
                meta = PAGES[index-1]
                shapes = list(walk(slide.shapes))
                text = "\n".join(s.text for s in shapes if s.has_text_frame)
                assert "SwiftFlow" in text and "HawkeyePlus" in text, "Joint branding missing"
                if lang == "en":
                    assert not re.search(r"[\u4e00-\u9fff]", text), (index, "Untranslated slide")
                    assert not re.search(r"[\u4e00-\u9fff]", slide.notes_slide.notes_text_frame.text), (index, "Untranslated notes")
                assert f"{index:02d} / {len(PAGES):02d}" in text, (index, "Incorrect page number")
                if meta["kind"] == "content":
                    chapter_index = meta["chapter"]
                    title = CHAPTERS[chapter_index]["title"][0 if lang == "zh" else 1]
                    kicker = (f"第 {chapter_index+1:02d} 章 / {title}" if lang == "zh"
                              else f"CHAPTER {chapter_index+1:02d} / {title.upper()}")
                    assert kicker in text, (index, "Wrong chapter label")
                if meta.get("content") == 9:
                    assert all(k in text for k in ["PP-YOLOE", "RT-DETR", "TorchDet", "FCOS", "SSD"])
                if meta.get("content") == 12:
                    assert ("Adaptation & joint acceptance" if lang == "en" else "适配与联合验收") in text
                if meta["kind"] == "cover":
                    assert ("Standardized Inspection & Measurement" if lang == "en" else "工业检测与测量标准化平台") in text
                stats = {
                    "text_boxes": sum(s.has_text_frame and bool(s.text.strip()) for s in shapes),
                    "pictures": sum(s.shape_type == MSO_SHAPE_TYPE.PICTURE for s in shapes),
                    "groups": sum(s.shape_type == MSO_SHAPE_TYPE.GROUP for s in shapes),
                    "native_objects": sum(s.shape_type != MSO_SHAPE_TYPE.PICTURE for s in shapes),
                }
                assert stats["text_boxes"] >= 8 and stats["pictures"] >= 2 and stats["groups"] >= 2, (index, stats)
                assert not any(s.shape_type == MSO_SHAPE_TYPE.PICTURE and
                               s.width > prs.slide_width*.9 and s.height > prs.slide_height*.9 for s in shapes)
                ids = [n.get("id") for n in slide._element.xpath(".//p:cNvPr")]
                assert len(ids) == len(set(ids))
                for s in slide.shapes:
                    assert s.left >= -12700 and s.top >= -12700, (index, s.name)
                    assert s.left+s.width <= prs.slide_width+12700, (index, s.name)
                    assert s.top+s.height <= prs.slide_height+12700, (index, s.name)
                record["slides"].append(stats)
            agenda_targets = {
                rel.target_part for rel in prs.slides[1].part.rels.values()
                if not rel.is_external and rel.reltype.endswith("/slide")
            }
            assert {prs.slides[c["first"]-1].part for c in CHAPTERS}.issubset(agenda_targets), "Broken agenda links"
            record["agenda_destinations"] = [c["first"] for c in CHAPTERS]
            presentation = app.Presentations.Open(str(path), ReadOnly=True, Untitled=False, WithWindow=False)
            try:
                sections = presentation.SectionProperties
                assert sections.Count == len(NATIVE_SECTIONS), "PowerPoint did not recognize native sections"
                record["native_sections"] = []
                for i, (first, names) in enumerate(NATIVE_SECTIONS, 1):
                    assert sections.FirstSlide(i) == first
                    assert sections.Name(i) == names[0 if lang == "zh" else 1]
                    record["native_sections"].append({"name": sections.Name(i), "first_slide": first})
                for index, meta in enumerate(PAGES, 1):
                    slide = presentation.Slides.Item(index)
                    for shape in com_walk(slide.Shapes):
                        if shape.HasTextFrame and shape.TextFrame.HasText:
                            tr = shape.TextFrame2.TextRange
                            if tr.BoundWidth > shape.Width+2 or tr.BoundHeight > shape.Height+2:
                                record["text_overflow"].append({
                                    "slide": index, "text": tr.Text, "bounds": [tr.BoundWidth, tr.BoundHeight],
                                    "shape": [shape.Width, shape.Height],
                                })
                    filename = meta["stem"]+"_"+suffix(lang, mode)+".png"
                    target = HERE/"verification"/filename
                    slide.Export(str(target), "PNG", 3840, 2160)
                    (HERE/"png"/filename).write_bytes(target.read_bytes())
                    with Image.open(target) as im:
                        assert im.size == (3840, 2160)
                presentation.SaveAs(str(HERE/f"{deck_stem(lang, mode)}.pdf"), 32)
            finally:
                presentation.Close()
            report.append(record)
    finally:
        # The COM server can share the user's PowerPoint process. Never Quit().
        pythoncom.CoUninitialize()
    (HERE/"verification/report.json").write_text(json.dumps({
        "powerpoint_opened": True, "slides_exported": len(PAGES)*len(LANGUAGES)*len(THEMES), "decks": report,
    }, ensure_ascii=False, indent=2), encoding="utf-8")
    assert not any(r["text_overflow"] for r in report), report
    contact_sheets()
    print(json.dumps([{
        "file": r["file"], "slides": len(r["slides"]), "native_sections": len(r["native_sections"]),
        "agenda_destinations": r["agenda_destinations"], "text_overflow": r["text_overflow"],
    } for r in report], indent=2))


if __name__ == "__main__":
    main()
